from __future__ import annotations

import csv
import hashlib
import io
import unicodedata
from pathlib import Path

from openpyxl import load_workbook


FIELDS = ("name", "precondition", "source_status", "priority", "step", "expected", "test_data", "folder")
LABELS = {
    "name": "Nome",
    "precondition": "Precondição",
    "source_status": "Status",
    "priority": "Prioridade",
    "step": "Passo",
    "expected": "Resultado Esperado",
    "test_data": "Dados do teste",
    "folder": "Pasta",
}
ALIASES = {
    "nome": "name",
    "precondicao": "precondition",
    "status": "source_status",
    "prioridade": "priority",
    "passo": "step",
    "resultado esperado": "expected",
    "dados do teste": "test_data",
    "pasta": "folder",
}
MAX_ROWS = 100_000


def normalized(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or "").strip().casefold())
    return " ".join("".join(ch for ch in text if not unicodedata.combining(ch)).split())


def as_text(value: object) -> str:
    return "" if value is None else str(value)


def file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _csv_sheets(path: Path):
    try:
        raw = path.read_text(encoding="utf-8-sig", errors="strict")
    except UnicodeDecodeError as exc:
        raise ValueError("O CSV precisa estar em UTF-8 (com ou sem BOM).") from exc
    if not raw.strip():
        raise ValueError("O CSV está vazio.")
    try:
        dialect = csv.Sniffer().sniff(raw[:8192], delimiters=",;")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ";" if raw.splitlines()[0].count(";") > raw.splitlines()[0].count(",") else ","
    try:
        reader = csv.reader(io.StringIO(raw, newline=""), delimiter=delimiter, strict=True)
        rows = list(reader)
    except csv.Error as exc:
        raise ValueError("O CSV está malformado. Confira aspas e separadores.") from exc
    if len(rows) > MAX_ROWS:
        raise ValueError(f"O limite é {MAX_ROWS:,} linhas por arquivo.")
    return [("CSV", rows)]


def _xlsx_sheets(path: Path):
    try:
        formulas = load_workbook(path, read_only=True, data_only=False)
        values = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("Não foi possível abrir o arquivo .xlsx. Verifique se não está corrompido.") from exc
    sheets = []
    try:
        for ws in formulas.worksheets:
            cached = values[ws.title]
            rows = []
            for row_number, (formula_row, cached_row) in enumerate(
                zip(ws.iter_rows(), cached.iter_rows()), start=1
            ):
                if row_number > MAX_ROWS:
                    raise ValueError(f"A aba {ws.title} excede {MAX_ROWS:,} linhas.")
                row_values = []
                for cell, cached_cell in zip(formula_row, cached_row):
                    if cell.data_type == "f":
                        row_values.append(cached_cell.value if cached_cell.value is not None else _Uncalculated(cell.coordinate))
                    else:
                        row_values.append(cell.value)
                rows.append(row_values)
            sheets.append((ws.title, rows))
    finally:
        formulas.close()
        values.close()
    return sheets


class _Uncalculated:
    def __init__(self, coordinate: str):
        self.coordinate = coordinate


def read_sheets(path: Path):
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return _csv_sheets(path)
    if suffix == ".xlsx":
        return _xlsx_sheets(path)
    raise ValueError("Envie um arquivo .xlsx ou .csv UTF-8.")


def _detect_map(headers: list[object], manual: dict[str, int] | None):
    if manual:
        result = {key: int(index) for key, index in manual.items() if key in FIELDS and index is not None}
        errors = [f"Mapeie a coluna obrigatória '{LABELS[key]}'." for key in ("name", "step") if key not in result]
        if len(set(result.values())) != len(result):
            errors.append("Cada campo deve usar uma coluna diferente.")
        if any(index < 0 or index >= len(headers) for index in result.values()):
            errors.append("O mapeamento aponta para uma coluna inexistente.")
        return result, errors
    result = {}
    errors = []
    normalized_headers = [normalized(header) for header in headers]
    if "passo" in normalized_headers and "acao" in normalized_headers:
        errors.append("Há colunas 'Passo' e 'Ação'. Confirme qual contém o texto da ação.")
    for index, header in enumerate(normalized_headers):
        field = ALIASES.get(header)
        if field:
            if field in result:
                errors.append(f"A coluna '{LABELS[field]}' aparece mais de uma vez.")
            else:
                result[field] = index
    for required in ("name", "step"):
        if required not in result:
            errors.append(f"A coluna obrigatória '{LABELS[required]}' não foi encontrada.")
    return result, errors


def inspect_file(path: Path, manual: dict[str, int] | None = None) -> dict:
    source_sheets = read_sheets(path)
    result = {
        "filename": path.name,
        "sha256": file_hash(path),
        "sheets": [],
        "errors": [],
        "total_cases": 0,
        "total_steps": 0,
        "headers": [],
        "mapping": {},
    }
    if not source_sheets:
        result["errors"].append("O arquivo não contém abas ou linhas.")
        return result
    first_headers = source_sheets[0][1][0] if source_sheets[0][1] else []
    result["headers"] = [as_text(v) for v in first_headers]
    field_map, map_errors = _detect_map(first_headers, manual)
    result["mapping"] = field_map
    result["errors"].extend(map_errors)
    if map_errors and ("name" not in field_map or "step" not in field_map or manual):
        return result
    for sheet_name, rows in source_sheets:
        sheet = {"name": sheet_name, "cases": [], "steps": 0, "errors": []}
        result["sheets"].append(sheet)
        if not rows:
            sheet["errors"].append("A aba está vazia.")
            continue
        if not manual:
            current_map, header_errors = _detect_map(rows[0], None)
            if header_errors:
                sheet["errors"].append("Cabeçalhos não reconhecidos nesta aba; revise o mapeamento.")
                continue
        else:
            current_map = field_map
        case = None
        for row_number, row in enumerate(rows[1:], start=2):
            if not any(value not in (None, "") for value in row):
                continue
            values = {}
            for field, index in current_map.items():
                value = row[index] if index < len(row) else None
                if isinstance(value, _Uncalculated):
                    sheet["errors"].append(f"Linha {row_number}, coluna {index + 1}: fórmula sem valor calculado.")
                    value = None
                values[field] = as_text(value)
            name = values.get("name", "").strip()
            step = values.get("step", "").strip()
            if name:
                case = {
                    "sheet": sheet_name,
                    "row": row_number,
                    "name": values["name"],
                    "precondition": values.get("precondition", ""),
                    "source_status": values.get("source_status", ""),
                    "priority": values.get("priority", ""),
                    "folder": values.get("folder", ""),
                    "steps": [],
                }
                sheet["cases"].append(case)
            elif case is None:
                sheet["errors"].append(f"Linha {row_number}: passo sem caso anterior.")
                continue
            elif any(values.get(field, "").strip() for field in ("precondition", "source_status", "priority", "folder")):
                sheet["errors"].append(f"Linha {row_number}: metadados do caso repetidos em linha de continuação.")
            if not step:
                sheet["errors"].append(f"Linha {row_number}: a coluna Passo está vazia.")
                continue
            case["steps"].append({
                "row": row_number,
                "action": values["step"],
                "expected": values.get("expected", ""),
                "test_data": values.get("test_data", ""),
            })
            sheet["steps"] += 1
        for item in sheet["cases"]:
            if not item["steps"]:
                sheet["errors"].append(f"Linha {item['row']}: caso sem passos.")
        result["total_cases"] += len(sheet["cases"])
        result["total_steps"] += sheet["steps"]
    return result
